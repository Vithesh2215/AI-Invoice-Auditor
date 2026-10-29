
from __future__ import annotations

import hashlib
import io
import logging
from pathlib import Path

from ..models import ExtractionResult
from ..settings import Settings

logger = logging.getLogger(__name__)

PDF_EXTENSIONS = {".pdf"}
DOCX_EXTENSIONS = {".docx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}


class ExtractionError(RuntimeError):
    """The document could not be read at all."""


def checksum_file(path: Path, chunk_size: int = 1 << 20) -> str:
    """Content checksum.

    Vendors rename files freely and the same invoice arrives twice under
    different names, so the filename is never the identity of an invoice.
    """
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(chunk_size), b""):
            digest.update(block)
    return digest.hexdigest()


def detect_format(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in PDF_EXTENSIONS:
        return "pdf"
    if suffix in DOCX_EXTENSIONS:
        return "docx"
    if suffix in IMAGE_EXTENSIONS:
        return suffix.lstrip(".")
    return suffix.lstrip(".") or "unknown"


def extract(path: Path, settings: Settings) -> ExtractionResult:
    suffix = path.suffix.lower()
    if suffix in PDF_EXTENSIONS:
        return _extract_pdf(path, settings)
    if suffix in DOCX_EXTENSIONS:
        return _extract_docx(path)
    if suffix in IMAGE_EXTENSIONS:
        return _extract_image(path, settings)
    raise ExtractionError(f"unsupported file type '{suffix}'")




def _extract_pdf(path: Path, settings: Settings) -> ExtractionResult:
    import pdfplumber

    warnings: list[str] = []
    pages: list[str] = []
    page_count = 0

    try:
        with pdfplumber.open(str(path)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except Exception as exc:
        warnings.append(f"pdfplumber failed: {type(exc).__name__}")
        logger.warning(
            "pdfplumber could not open file",
            extra={"file": path.name, "error": type(exc).__name__},
        )

    text = "\n".join(pages).strip()

    if len(text) >= settings.ocr_min_text_chars:
        return ExtractionResult(
            raw_text=text,
            extraction_method="pdfplumber",
            source_format="pdf",
            page_count=page_count or 1,
            ocr_used=False,
            warnings=warnings,
        )

    warnings.append("no usable embedded text layer; falling back to OCR")
    ocr_text, ocr_pages = _ocr_pdf(path, settings)
    if not ocr_text.strip():
        raise ExtractionError(
            "PDF has no embedded text and OCR produced nothing readable"
        )

    return ExtractionResult(
        raw_text=ocr_text,
        extraction_method="pdf2image+tesseract",
        source_format="pdf",
        page_count=ocr_pages or page_count or 1,
        ocr_used=True,
        warnings=warnings,
    )


def _ocr_pdf(path: Path, settings: Settings) -> tuple[str, int]:
    import pytesseract
    from pdf2image import convert_from_path

    try:
        images = convert_from_path(str(path), dpi=settings.ocr_dpi)
    except Exception as exc:
        raise ExtractionError(
            f"could not rasterise PDF for OCR ({type(exc).__name__}); "
            "is poppler-utils installed in the image?"
        ) from None

    texts = [
        pytesseract.image_to_string(image, lang=settings.ocr_languages)
        for image in images
    ]
    return "\n".join(texts).strip(), len(images)




def _extract_docx(path: Path) -> ExtractionResult:
    from docx import Document

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"could not open DOCX ({type(exc).__name__})") from None

    lines = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    text = "\n".join(lines).strip()
    if not text:
        raise ExtractionError("DOCX contained no readable text")

    return ExtractionResult(
        raw_text=text,
        extraction_method="python-docx",
        source_format="docx",
        page_count=1,
        ocr_used=False,
    )




def _extract_image(path: Path, settings: Settings) -> ExtractionResult:
    import pytesseract
    from PIL import Image, ImageOps

    try:
        with Image.open(path) as image:
            prepared = _prepare_for_ocr(image, ImageOps)
            text = pytesseract.image_to_string(prepared, lang=settings.ocr_languages)
    except Exception as exc:
        raise ExtractionError(
            f"OCR failed for image ({type(exc).__name__}); "
            "is tesseract installed with the required language packs?"
        ) from None

    if not text.strip():
        raise ExtractionError("OCR produced no readable text from the image")

    return ExtractionResult(
        raw_text=text.strip(),
        extraction_method="tesseract",
        source_format=path.suffix.lower().lstrip("."),
        page_count=1,
        ocr_used=True,
    )


def _prepare_for_ocr(image, image_ops):
    """Greyscale, autocontrast and upscale small scans before OCR."""
    prepared = image.convert("L")
    prepared = image_ops.autocontrast(prepared)
    if min(prepared.size) < 1000:
        scale = 1000 / min(prepared.size)
        prepared = prepared.resize(
            (int(prepared.width * scale), int(prepared.height * scale))
        )
    buffer = io.BytesIO()
    prepared.save(buffer, format="PNG")
    buffer.seek(0)
    from PIL import Image as PILImage

    return PILImage.open(buffer)
